# -*- coding: utf-8 -*-
"""
Markdown → DOCX 转换器（面向本项目的交付文档）

覆盖范围：标题、段落、无序/有序列表、表格、引用块、代码块、分隔线、
行内 **粗体** 与 `代码`。未支持的语法按普通段落输出，不会丢失内容。

中文字体处理：python-docx 默认只设 latin 字形，中文会回退成宋体/等线。
这里显式写入 w:eastAsia，保证中英文都用指定字体渲染。

用法：
    python deliverables/build/md_to_docx.py <input.md> <output.docx>
"""
from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BODY_FONT = "微软雅黑"
MONO_FONT = "Consolas"
INK = RGBColor(0x1A, 0x1A, 0x1A)
INK_SOFT = RGBColor(0x55, 0x55, 0x55)
ACCENT = RGBColor(0x1F, 0x5C, 0xA8)


def _style_font(run, name: str = BODY_FONT) -> None:
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


_INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text: str, size: int = 10.5, base_bold: bool = False,
               color: RGBColor = INK) -> None:
    """写入一段可能含 **粗体** / `代码` 的文本。"""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    text = text.replace("\\|", "|")
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            _style_font(run)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            _style_font(run, MONO_FONT)
            run.font.size = Pt(size - 0.5)
            run.font.color.rgb = ACCENT
        else:
            run = paragraph.add_run(part)
            run.bold = base_bold
            _style_font(run)
        run.font.size = Pt(size)
        if not (part.startswith("`") and part.endswith("`")):
            run.font.color.rgb = color


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip()))


def convert(md_path: str, docx_path: str, title: str | None = None) -> None:
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    # 页边距：给表格留出宽度
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(60)
        section.top_margin = section.bottom_margin = Pt(60)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    i = 0
    in_code = False
    code_buf: list[str] = []

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        # ── 代码块 ──
        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(14)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_buf))
                _style_font(run, MONO_FONT)
                run.font.size = Pt(9)
                run.font.color.rgb = INK_SOFT
                code_buf, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        stripped = line.strip()

        # ── 空行 / 分隔线 ──
        if not stripped:
            i += 1
            continue
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("─" * 46)
            _style_font(run)
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            i += 1
            continue

        # ── 表格 ──
        if stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            header = split_row(stripped)
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(split_row(lines[j]))
                j += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ci, text in enumerate(header):
                cell = table.rows[0].cells[ci]
                cell.text = ""
                add_inline(cell.paragraphs[0], text, size=9.5, base_bold=True)
            for row in rows:
                cells = table.add_row().cells
                for ci in range(len(header)):
                    cell = cells[ci]
                    cell.text = ""
                    add_inline(cell.paragraphs[0], row[ci] if ci < len(row) else "", size=9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i = j
            continue

        # ── 标题 ──
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            p.paragraph_format.space_after = Pt(6)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, m.group(2), size=18, base_bold=True)
            elif level == 2:
                add_inline(p, m.group(2), size=14.5, base_bold=True)
            elif level == 3:
                add_inline(p, m.group(2), size=12.5, base_bold=True)
            else:
                add_inline(p, m.group(2), size=11, base_bold=True)
            i += 1
            continue

        # ── 引用块 ──
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(16)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, stripped.lstrip(">").strip(), size=10, color=INK_SOFT)
            i += 1
            continue

        # ── 列表 ──
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", raw)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Pt(18 + 14 * (len(m.group(1)) // 2))
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", raw)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Pt(18 + 14 * (len(m.group(1)) // 2))
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, m.group(3))
            i += 1
            continue

        # ── 普通段落 ──
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.35
        add_inline(p, stripped)
        i += 1

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    convert(sys.argv[1], sys.argv[2])
    print(f"OK {sys.argv[2]}")
